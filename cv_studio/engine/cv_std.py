# -*- coding: utf-8 -*-
"""Standardise Zakheni candidate CVs from a dropped-in folder or zip.

Engine used by both the CLI and the web UI.

CLI:
  python cv_std.py unpack <input.zip|dir>          ingest a pack into %TEMP%\\opencode\\cvstd\\<id>
  python cv_std.py inspect <pack_id>               show what was found
  python cv_std.py skeleton <pack_id> [--role X]   scaffold candidates_<pack_id>.py to fill in
  python cv_std.py build <pack_id> [--dest DIR]    generate DOCX+PDF, copy docs, deploy + verify
  python cv_std.py verify <pack_id> [--dest DIR]   check a delivered pack
"""
import argparse
import glob
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

ROOT_TEMP = os.environ.get("CVSTD_WORK_ROOT") or os.path.join(
    os.environ.get("TEMP", tempfile.gettempdir()), "opencode", "cvstd"
)
DEST_ROOT = os.environ.get("CVSTD_DEST_ROOT") or \
    r"C:\Users\jerma\OneDrive - Zakheni ICT (Pty) Ltd\Desktop\CVs\Standardised CVs"


def set_roots(work_root=None, dest_root=None):
    """Override the engine's work/destination roots at runtime (used by the
    Odoo module, which cannot rely on import-time environment variables)."""
    global ROOT_TEMP, DEST_ROOT
    if work_root:
        ROOT_TEMP = work_root
    if dest_root:
        DEST_ROOT = dest_root

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import gen_cv
from pdf_conv import convert as docx_to_pdf

CV_RE = re.compile(r"(?i)(^|[\s._-])(cv|curriculum[\s._-]*vitae|resume)([\s._-]|$)")
DOCDIR_RE = re.compile(r"(?i)qual|qal|cert|credential|proof|document|degree|diplomac|academic")
SCOPE_RE = re.compile(r"(?i)scope|statement\s*of\s*work|sow|terms?\b|brief|requirement")
DOC_EXT = {".pdf", ".docx", ".doc", ".txt"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff"}
SHEET_EXT = {".xlsx", ".xls"}

BASE_RE = re.compile(
    r"(?i)(qualif|credential|proof|degree|baccalaure|diploma|matric|"
    r"certif|academic|transcript|award|identity|national|"
    r"\bbsc\b|\bbcom\b|\bmba\b|(^|[^a-z])id([^a-z]|$))"
)


# ---------------------------------------------------------------- helpers

def slug(s):
    s = re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")
    return re.sub(r"_+", "_", s)


def work_dir(pack_id):
    return os.path.join(ROOT_TEMP, pack_id)


def clean_name(n):
    return re.sub(r"\s+", " ", n).strip().strip("_").strip("-")


def load_manifest(pack_id):
    p = os.path.join(work_dir(pack_id), "manifest.json")
    if not os.path.exists(p):
        raise FileNotFoundError(f"no manifest for pack '{pack_id}' (run unpack first)")
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def save_manifest(pack_id, mf):
    p = os.path.join(work_dir(pack_id), "manifest.json")
    with open(p, "w", encoding="utf-8") as f:
        json.dump(mf, f, indent=2, ensure_ascii=False)
    return p


def rel(base, full):
    return os.path.relpath(full, base)


def find_common_top(paths):
    tops = set()
    for p in paths:
        parts = [x for x in p.replace("\\", "/").split("/") if x]
        if parts:
            tops.add(parts[0])
    if len(tops) == 1:
        top = next(iter(tops))
        if all(p.replace("\\", "/").startswith(top + "/") for p in paths):
            return top
    return None


# ---------------------------------------------------------------- extraction

def extract_docx(path):
    """Extract text from a DOCX in document order, including table cells.

    CVs built from templates often put all content inside tables, which the
    plain `doc.paragraphs` view misses entirely.
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = Document(path)
    lines = []

    def cell_text(cell):
        parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
        for t in cell.tables:
            for row in t.rows:
                sub = [cell_text(c) for c in row.cells]
                parts.append(" | ".join(s for s in sub if s))
        return " ".join(parts)

    def walk(table):
        for row in table.rows:
            cells = [cell_text(c) for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line.strip():
                lines.append(line)

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, d).text
            if t.strip():
                lines.append(t)
        elif tag == "tbl":
            walk(Table(child, d))
    return "\n".join(lines)


def extract_pdf(path):
    try:
        import pymupdf
        doc = pymupdf.open(path)
        return "\n".join(pg.get_text() for pg in doc), doc.page_count
    except Exception:
        import fitz
        doc = fitz.open(path)
        return "\n".join(pg.get_text() for pg in doc), doc.page_count


def convert_doc(path, outroot):
    os.makedirs(outroot, exist_ok=True)
    out = os.path.join(outroot, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    if os.path.exists(out):
        return out
    import subprocess, time
    import win32com.client as win32
    try:
        import pythoncom
        pythoncom.CoInitialize()
    except Exception:
        pass
    subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
    time.sleep(2)
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(path, ReadOnly=True)
        doc.SaveAs2(out, FileFormat=17)
        doc.Close(False)
    finally:
        try:
            word.Quit()
        except Exception:
            pass
        subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
        time.sleep(1)
    return out


def extract_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"=== Sheet: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in row]
            if any(vals):
                lines.append(" | ".join(vals))
    return "\n".join(lines)


def render_pdf_pages(pdf, outdir, dpi=300):
    import pymupdf
    os.makedirs(outdir, exist_ok=True)
    out = []
    doc = pymupdf.open(pdf)
    for i, pg in enumerate(doc):
        p = os.path.join(outdir, f"p{i}.png")
        pg.get_pixmap(dpi=dpi).save(p)
        out.append(p)
    return out


def render_image(img, outdir):
    from PIL import Image
    os.makedirs(outdir, exist_ok=True)
    p = os.path.join(outdir, "img.png")
    Image.open(img).convert("RGB").save(p)
    return [p]


# ---------------------------------------------------------------- unpack

def classify_files(cand_root, wdir):
    """Split a candidate's files into cvs vs docs using heuristics."""
    cvs, docs, unclassified = [], [], []
    for root, dirs, files in os.walk(cand_root):
        sub = os.path.relpath(root, cand_root)
        in_docdir = sub != "." and bool(DOCDIR_RE.search(sub))
        for fn in sorted(files):
            full = os.path.join(root, fn)
            ext = os.path.splitext(fn)[1].lower()
            if ext not in DOC_EXT and ext not in IMG_EXT:
                continue
            entry = {"file": rel(wdir, full), "base": fn, "ext": ext}
            if in_docdir:
                docs.append(entry)
            elif ext in DOC_EXT and CV_RE.search(fn):
                cvs.append(entry)
            elif BASE_RE.search(fn):
                docs.append(entry)
            else:
                unclassified.append(entry)
    if not cvs and len(unclassified) == 1 and unclassified[0]["ext"] in DOC_EXT:
        cvs.append(unclassified.pop())
    return cvs, docs, unclassified


def has_docs(dirpath):
    """True if the subtree contains at least one document or image (candidate media)."""
    for r, _, fs in os.walk(dirpath):
        for fn in fs:
            if os.path.splitext(fn)[1].lower() in (DOC_EXT | IMG_EXT):
                return True
    return False


def is_person_dir(dirpath):
    """A role-folder child that is a person folder, not a docs/qualifications folder."""
    name = os.path.basename(dirpath.rstrip("\\/"))
    if DOCDIR_RE.search(name):
        return False
    return has_docs(dirpath)


def unpack(input_path, pack_id=None, no_ocr=False):
    if not os.path.exists(input_path):
        raise FileNotFoundError(input_path)
    pack_id = pack_id or slug(os.path.splitext(os.path.basename(input_path))[0])
    wdir = work_dir(pack_id)
    srcdir = os.path.join(wdir, "src")
    if os.path.isdir(wdir):
        shutil.rmtree(wdir)
    os.makedirs(srcdir, exist_ok=True)

    if os.path.isdir(input_path):
        root = os.path.join(srcdir, "input")
        shutil.copytree(input_path, root)
    else:
        with zipfile.ZipFile(input_path) as z:
            names = z.namelist()
            common = find_common_top(names)
            root = os.path.join(srcdir, common) if common else srcdir
            for n in names:
                if n.endswith("/"):
                    continue
                dest = os.path.join(srcdir, n)
                os.makedirs(os.path.dirname(dest), exist_ok=True)
                with z.open(n) as src, open(dest, "wb") as out:
                    shutil.copyfileobj(src, out)

    mf = {
        "pack_id": pack_id,
        "pack_name": clean_name(os.path.splitext(os.path.basename(input_path))[0]),
        "input": input_path,
        "scope": [],
        "spreadsheet": [],
        "candidates": [],
    }

    entries = sorted(os.listdir(root))
    subdirs = [e for e in entries if os.path.isdir(os.path.join(root, e))]
    direct_files = [e for e in entries if os.path.isfile(os.path.join(root, e))]

    def add_pack_level_file(full):
        ext = os.path.splitext(full)[1].lower()
        name = os.path.basename(full)
        target = None
        if SCOPE_RE.search(name) and ext in DOC_EXT:
            target = "scope"
        elif ext in SHEET_EXT or name.lower().endswith(".xls") or "spreadsheet" in name.lower():
            target = "spreadsheet"
        if target:
            txt = extract_any(full, os.path.join(wdir, "txt", "pack"))
            mf[target].append({"file": rel(wdir, full), "text_txt": rel(wdir, txt) if txt else None})

    for f in direct_files:
        add_pack_level_file(os.path.join(root, f))

    cand_folders = [os.path.join(root, d) for d in subdirs
                    if any(os.path.isfile(os.path.join(r, x))
                           for r, _, fs in os.walk(os.path.join(root, d)) for x in fs)]

    seen_keys = set()

    def make_key(name, role):
        k = slug(name) or "candidate"
        if k in seen_keys and role:
            k = slug(role + "_" + name) or k
        n = 2
        while k in seen_keys:
            k = "%s_%d" % (slug(name) or "candidate", n)
            n += 1
        seen_keys.add(k)
        return k

    def ingest(cand_folder, folder_name, role=None):
        cvs, docs, uncl = classify_files(cand_folder, wdir)
        cand = {
            "key": make_key(clean_name(os.path.basename(folder_name)), role),
            "name": clean_name(os.path.basename(folder_name)),
            "folder": os.path.basename(folder_name),
            "role": role,
            "cvs": [],
            "docs": [],
            "unclassified": [{"file": u["file"], "base": u["base"]} for u in uncl],
        }
        for c in cvs:
            txt = extract_any(
                os.path.join(wdir, c["file"]), os.path.join(wdir, "txt", cand["key"])
            )
            cand["cvs"].append({"file": c["file"], "base": c["base"],
                                "text_txt": rel(wdir, txt) if txt else None})
        for d in docs:
            full = os.path.join(wdir, d["file"])
            txt_path = extract_any(full, os.path.join(wdir, "txt", cand["key"]))
            textable = False
            if txt_path:
                with open(txt_path, encoding="utf-8") as fh:
                    textable = len(fh.read().strip()) > 40
            ocr = []
            if not no_ocr and (not textable or d["ext"] in IMG_EXT):
                od = os.path.join(wdir, "ocr", cand["key"], slug(d["base"]))
                try:
                    if d["ext"] in DOC_EXT and d["ext"] == ".pdf":
                        ocr = render_pdf_pages(full, od)
                    elif d["ext"] in IMG_EXT:
                        ocr = render_image(full, od)
                    else:
                        ocr = []
                except Exception:
                    ocr = []
            cand["docs"].append({
                "file": d["file"], "base": d["base"],
                "text_txt": rel(wdir, txt_path) if txt_path else None,
                "textable": textable, "pages": None, "ocr": [rel(wdir, x) for x in ocr],
            })
            if d["ext"] == ".pdf":
                try:
                    cand["docs"][-1]["pages"] = page_count(full)
                except Exception:
                    pass
        mf["candidates"].append(cand)

    if cand_folders:
        for cf in cand_folders:
            top = os.path.basename(cf)
            person_dirs = [os.path.join(cf, cd) for cd in sorted(os.listdir(cf))
                           if os.path.isdir(os.path.join(cf, cd)) and is_person_dir(os.path.join(cf, cd))]
            if person_dirs:
                for pd in person_dirs:
                    ingest(pd, os.path.basename(pd), role=clean_name(top))
            else:
                ingest(cf, top, role=None)
    else:
        ingest(root, os.path.basename(root) or mf["pack_name"], role=None)

    save_manifest(pack_id, mf)
    return mf


def page_count(pdf):
    import pymupdf
    return pymupdf.open(pdf).page_count


def extract_any(full, outdir):
    ext = os.path.splitext(full)[1].lower()
    os.makedirs(outdir, exist_ok=True)
    if ext == ".docx":
        txt = extract_docx(full)
    elif ext == ".pdf":
        txt, _ = extract_pdf(full)
    elif ext == ".doc":
        conv = os.path.join(os.path.dirname(os.path.dirname(outdir)), "_converted")
        pdf = convert_doc(full, conv)
        txt, _ = extract_pdf(pdf)
    elif ext == ".txt":
        with open(full, encoding="utf-8", errors="replace") as fh:
            txt = fh.read()
    elif ext in SHEET_EXT:
        txt = extract_xlsx(full)
    else:
        return None
    base = os.path.splitext(os.path.basename(full))[0]
    outp = os.path.join(outdir, base + ".txt")
    n = 1
    while os.path.exists(outp):
        outp = os.path.join(outdir, f"{base}_{n}.txt")
        n += 1
    with open(outp, "w", encoding="utf-8") as fh:
        fh.write(txt)
    return outp


def reextract(pack_id):
    """Re-run text extraction for every file in a pack, overwriting the
    existing txt files in place (e.g. after improving the extractor)."""
    mf = load_manifest(pack_id)
    wdir = work_dir(pack_id)
    n = 0
    for c in mf["candidates"]:
        for item in list(c.get("cvs", [])) + list(c.get("docs", [])):
            tt = item.get("text_txt")
            full = os.path.join(wdir, item.get("file", ""))
            if not tt or not os.path.exists(full):
                continue
            ext = os.path.splitext(full)[1].lower()
            if ext == ".docx":
                txt = extract_docx(full)
            elif ext == ".pdf":
                txt, _ = extract_pdf(full)
            elif ext == ".txt":
                with open(full, encoding="utf-8", errors="replace") as fh:
                    txt = fh.read()
            elif ext in SHEET_EXT:
                txt = extract_xlsx(full)
            else:
                continue
            with open(os.path.join(wdir, tt), "w", encoding="utf-8") as fh:
                fh.write(txt)
            n += 1
    return {"reextracted": n}


# ---------------------------------------------------------------- skeleton

def skeleton(pack_id, role=None, force=False):
    mf = load_manifest(pack_id)
    outp = os.path.join(work_dir(pack_id), f"candidates_{pack_id}.py")
    if os.path.exists(outp) and not force:
        return {"created": False, "saved": False, "path": outp,
                "msg": "candidates file already exists (use --force to overwrite)"}
    lines = ["# -*- coding: utf-8 -*-", "CANDIDATES = {"]
    for c in mf["candidates"]:
        key = c["key"]
        lines.append(f'  "{key}": {{')
        lines.append(f'    "name": {json.dumps(c["name"], ensure_ascii=False)},')
        lines.append(f'    "role": {json.dumps(c.get("role") or "", ensure_ascii=False)},')
        lines.append(f'    "target_role": {json.dumps(role or c.get("role") or "", ensure_ascii=False)},')
        lines.append(f'    "summary": [{json.dumps("", ensure_ascii=False)}],')
        lines.append('    "competencies": [],')
        lines.append('    "certs": [],')
        lines.append('    "education": [],')
        lines.append('    "career": [],')
        lines.append('    "refs": ["Available on request."],')
        docs = ['      ' + json.dumps(f"<describe this document> (see: {cand_doc['base']})",
                                        ensure_ascii=False) + ","
                for cand_doc in c["docs"]]
        if docs:
            lines.append("    \"docs\": [")
            lines.extend(docs)
            lines.append("    ],")
        else:
            lines.append('    "docs": [],')
        lines.append("  },")
    lines.append("}")
    with open(outp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return {"created": True, "saved": True, "path": outp}


def save_candidates(pack_id, candidates):
    outp = os.path.join(work_dir(pack_id), f"candidates_{pack_id}.py")
    lines = ["# -*- coding: utf-8 -*-", "CANDIDATES = {"]
    for key, c in candidates.items():
        key = slug(str(key)) or "candidate"
        lines.append(f'  "{key}": {json.dumps(c, ensure_ascii=False, indent=2)},'.replace("\n", "\n  "))
    lines.append("}")
    with open(outp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return outp


def load_candidates(pack_id):
    p = os.path.join(work_dir(pack_id), f"candidates_{pack_id}.py")
    if not os.path.exists(p):
        raise FileNotFoundError(p)
    ns = {}
    with open(p, encoding="utf-8") as fh:
        code = fh.read()
    exec(compile(code, p, "exec"), ns)
    return ns["CANDIDATES"]


def primary_cv_text(pack_id, cand):
    best = ""
    for cv in cand.get("cvs", []):
        if cv.get("text_txt"):
            p = os.path.join(work_dir(pack_id), cv["text_txt"])
            if os.path.exists(p):
                with open(p, encoding="utf-8", errors="replace") as fh:
                    txt = fh.read()
                if len(txt.strip()) > len(best.strip()):
                    best = txt
    if best.strip():
        return best
    for d in cand.get("docs", []):
        if d.get("text_txt") and d.get("textable"):
            p = os.path.join(work_dir(pack_id), d["text_txt"])
            if os.path.exists(p):
                with open(p, encoding="utf-8", errors="replace") as fh:
                    txt = fh.read()
                if len(txt.strip()) > len(best.strip()):
                    best = txt
    return best


def primary_cv_path(pack_id, cand, cvs_only=False):
    """Path to the best extracted CV text file for a candidate (or None).

    cvs_only=True restricts to real CV files (ignores supporting docs), so the
    caller can tell whether the person actually has usable CV text.
    """
    best, best_len = None, -1
    for cv in cand.get("cvs", []):
        if cv.get("text_txt"):
            p = os.path.join(work_dir(pack_id), cv["text_txt"])
            if os.path.exists(p):
                n = os.path.getsize(p)
                if n > best_len:
                    best, best_len = p, n
    if best or cvs_only:
        return best
    for d in cand.get("docs", []):
        if d.get("text_txt") and d.get("textable"):
            p = os.path.join(work_dir(pack_id), d["text_txt"])
            if os.path.exists(p):
                n = os.path.getsize(p)
                if n > best_len:
                    best, best_len = p, n
    return best


def draft(pack_id, force=False, llm=False, model=None, progress=None,
          retries=1, timeout=240, only=None, resume=False, min_cv_chars=1500,
          max_cv_chars=20000):
    """Auto-fill candidates from each person's source CV text.

    llm=False -> fast heuristic parser (rusty quality).
    llm=True  -> ask opencode.exe to curate each CV (house quality, slow).
    resume=True -> keep candidates already drafted, only fill the rest.
    Candidates whose real CV text is missing/tiny are never sent to the LLM
    (it would stall and hallucinate); they fall back to the heuristic parser
    and are flagged with _weak.
    """
    mf = load_manifest(pack_id)
    outp = os.path.join(work_dir(pack_id), f"candidates_{pack_id}.py")
    exists = os.path.exists(outp)
    if exists and not force and not resume:
        return {"created": False, "path": outp, "msg": "exists (use --force)"}
    C = {}
    if resume and exists:
        try:
            C = load_candidates(pack_id)
        except Exception:  # noqa: BLE001
            C = {}
    empty = 0
    weak = 0
    failed = 0
    only_set = set(only) if only else set()
    cands = mf["candidates"]
    if only_set:
        cands = [c for c in cands if c["key"] in only_set]
    total = len(cands)
    for i, c in enumerate(cands):
        docs = c.get("docs", [])
        if progress:
            progress(i + 1, total, c["name"])
        if (c["key"] in C and C[c["key"]].get("target_role", "").strip()
                and c["key"] not in only_set):
            continue
        cv_path = primary_cv_path(pack_id, c, cvs_only=True)
        has_cv = bool(cv_path) and os.path.getsize(cv_path) >= min_cv_chars
        if llm and has_cv:
            import llm_draft
            got = None
            last_err = None
            for _ in range(retries + 1):
                try:
                    got = llm_draft.draft_via_llm(
                        cv_path, name=c["name"], role=c.get("role") or "",
                        docs=docs, model=model, timeout=timeout,
                        max_chars=max_cv_chars)
                    break
                except Exception as e:  # timeout / bad JSON / crash
                    last_err = e
            if got is None:
                failed += 1
                got = None
        else:
            got = None
        if got is None:
            import autodraft
            weak += 1
            got = autodraft.draft_from_text(
                primary_cv_text(pack_id, c), name=c["name"],
                role=c.get("role"), docs=docs)
            if not has_cv:
                empty += 1
                got["_weak"] = "no usable CV text (scanned or image CV)"
            elif llm:
                got["_weak"] = "LLM draft failed; fast draft used"
        C[c["key"]] = got
        save_candidates(pack_id, C)
    save_candidates(pack_id, C)
    return {"created": True, "path": outp, "candidates": len(C),
            "no_text": empty, "llm": llm, "llm_failed": failed, "weak": weak}


# ---------------------------------------------------------------- build / verify

def resolve_doc_path(pack_id, base):
    wdir = work_dir(pack_id)
    for p in glob.glob(os.path.join(wdir, "src", "**", "*"), recursive=True):
        if os.path.basename(p) == base:
            return p
    converted = os.path.join(wdir, "_converted")
    if os.path.isdir(converted):
        for p in glob.glob(os.path.join(converted, "**", "*"), recursive=True):
            if os.path.basename(p) == base:
                return p
    return None


def build(pack_id, dest=None, no_deploy=False, skip_verify=False,
          skip_weak=False):
    mf = load_manifest(pack_id)
    C = load_candidates(pack_id)
    wdir = work_dir(pack_id)
    dest = dest or os.path.join(DEST_ROOT, mf["pack_name"])
    skipped = []
    if skip_weak:
        for k in list(C):
            if C[k].get("_weak"):
                skipped.append(C[k].get("name") or k)
                del C[k]

    problems = []
    for key, c in C.items():
        if not c.get("target_role", "").strip():
            problems.append(f"{c.get('name', key)}: target_role is empty")

    deploy = os.path.join(wdir, "deploy", mf["pack_name"])
    if os.path.isdir(deploy):
        shutil.rmtree(deploy)
    os.makedirs(deploy, exist_ok=True)

    for key, c in C.items():
        c.setdefault("docs", [])
        gen_cv.build(c, os.path.join(deploy, gen_cv.to_filename(c["name"]) + ".docx"))

    docx_to_pdf(deploy, deploy)

    report = {"deployed": [], "docs_copied": 0, "missing_refs": [], "problems": problems}

    for key, c in C.items():
        name = c["name"]
        if no_deploy:
            continue
        role = (c.get("role") or "").strip()
        cand_dst = os.path.join(dest, role, name) if role else os.path.join(dest, name)
        os.makedirs(cand_dst, exist_ok=True)
        for f in [gen_cv.to_filename(name) + ".docx", gen_cv.to_filename(name) + ".pdf"]:
            s = os.path.join(deploy, f)
            if os.path.exists(s):
                shutil.copy2(s, os.path.join(cand_dst, f))
                report["deployed"].append(f)
        for d in c.get("docs", []):
            m = re.search(r"\(see: (.+)\)", d)
            if not m:
                continue
            fname = m.group(1).strip()
            src = resolve_doc_path(pack_id, fname)
            if not src:
                report["missing_refs"].append((name, fname))
                continue
            shutil.copy2(src, os.path.join(cand_dst, fname))
            report["docs_copied"] += 1

    if not skip_verify:
        report["verify"] = verify_pack(dest, mf["pack_name"])
    if skipped:
        report["skipped_weak"] = skipped
    return report


KNOWN_HEADINGS = ["PROFESSIONAL SUMMARY", "CORE TECHNICAL COMPETENCIES", "CERTIFICATIONS",
                  "EDUCATION", "PROFESSIONAL EXPERIENCE", "REFERENCES", "SUPPORTING DOCUMENTS"]


def verify_pack(dest, pack_name):
    from docx import Document
    results = []
    cand_dirs = set()
    for r, _, fs in os.walk(dest):
        if any(f.endswith(" - Zakheni CV.docx") for f in fs):
            cand_dirs.add(r)
    for cdir in sorted(cand_dirs):
        d = os.path.basename(cdir)
        docx = os.path.join(cdir, d + " - Zakheni CV.docx")
        pdf = os.path.join(cdir, d + " - Zakheni CV.pdf")
        entry = {"candidate": d, "ok": True, "checks": []}
        if not os.path.exists(docx) or not os.path.exists(pdf):
            entry["ok"] = False
            entry["checks"].append(f"missing pair: docx={os.path.exists(docx)} pdf={os.path.exists(pdf)}")
            results.append(entry)
            continue
        doc = Document(docx)
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        heads = [t for t in texts if t.upper() in KNOWN_HEADINGS]
        idxs = [texts.index(h) for h in heads]
        if idxs != sorted(idxs):
            entry["ok"] = False
            entry["checks"].append("headings out of order")
        role = any(t.startswith("Target Role:") and t != "Target Role:" for t in texts)
        if not role:
            entry["ok"] = False
            entry["checks"].append("no Target Role line")
        have = set(os.listdir(cdir))
        refs = [m.group(1) for t in texts for m in [re.search(r"\(see: (.+)\)", t)] if m]
        missing = [r for r in refs if r not in have]
        if missing:
            entry["ok"] = False
            entry["checks"].append(f"unresolved refs: {missing}")
        entry["checks"].append(f"sections: {len(heads)}, refs: {len(refs)}")
        results.append(entry)
    return {"pack": pack_name, "candidates": results}


# ---------------------------------------------------------------- CLI

def main(argv=None):
    ap = argparse.ArgumentParser(description="Zakheni CV standardiser")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_unpack = sub.add_parser("unpack", help="ingest a zip or folder")
    p_unpack.add_argument("input")
    p_unpack.add_argument("--id", default=None)
    p_unpack.add_argument("--no-ocr", action="store_true")

    p_inspect = sub.add_parser("inspect", help="show manifest summary")
    p_inspect.add_argument("pack_id")

    p_reex = sub.add_parser("reextract",
                            help="re-extract text for all files in a pack")
    p_reex.add_argument("pack_id")

    p_skeleton = sub.add_parser("skeleton", help="scaffold candidates file")
    p_skeleton.add_argument("pack_id")
    p_skeleton.add_argument("--role", default=None)
    p_skeleton.add_argument("--force", action="store_true")

    p_draft = sub.add_parser("draft", help="auto-fill candidates from source CV text")
    p_draft.add_argument("pack_id")
    p_draft.add_argument("--force", action="store_true")
    p_draft.add_argument("--resume", action="store_true",
                         help="keep already-drafted candidates, fill the rest")
    p_draft.add_argument("--llm", action="store_true",
                         help="use opencode/LLM to curate (house quality, slow)")
    p_draft.add_argument("--model", default=None,
                         help="provider/model for --llm (default opencode/big-pickle)")
    p_draft.add_argument("--only", default=None,
                         help="comma-separated candidate keys to (re)draft")
    p_draft.add_argument("--max-chars", type=int, default=20000,
                         help="max CV text sent to the LLM (head+tail kept)")

    p_build = sub.add_parser("build", help="generate + deploy + verify")
    p_build.add_argument("pack_id")
    p_build.add_argument("--dest", default=None)
    p_build.add_argument("--no-deploy", action="store_true")
    p_build.add_argument("--skip-verify", action="store_true")
    p_build.add_argument("--skip-weak", action="store_true",
                         help="exclude candidates flagged _weak (no usable CV)")

    p_verify = sub.add_parser("verify", help="verify a delivered pack")
    p_verify.add_argument("pack_id")
    p_verify.add_argument("--dest", default=None)

    a = ap.parse_args(argv)
    if a.cmd == "unpack":
        mf = unpack(a.input, a.id, no_ocr=a.no_ocr)
        print(f"pack_id        : {mf['pack_id']}")
        print(f"pack_name      : {mf['pack_name']}")
        print(f"work dir       : {work_dir(mf['pack_id'])}")
        for c in mf["candidates"]:
            print(f"  {c['name']:<30} cvs={len(c['cvs'])} docs={len(c['docs'])} uncl={len(c['unclassified'])}")
        if not a.no_ocr:
            print("note: OCR rendered for scanned/multi-page supporting docs")
        print("drafted:", draft(mf["pack_id"]))
    elif a.cmd == "inspect":
        mf = load_manifest(a.pack_id)
        print(f"pack: {mf['pack_name']}  ({mf['pack_id']})")
        for c in mf["candidates"]:
            print(f"\n== {c['name']}  [{c['key']}]")
            for cv in c["cvs"]:
                print(f"   CV   {cv['file']}" + ("  (no text)" if not cv["text_txt"] else ""))
            for d in c["docs"]:
                tag = "" if d["textable"] else "  [scan -> ocr]"
                print(f"   DOC  {d['file']}{tag}")
            for u in c["unclassified"]:
                print(f"   ???  {u['file']}")
        for s in mf["scope"]:
            print(f"scope: {s['file']}")
        for s in mf["spreadsheet"]:
            print(f"spreadsheet: {s['file']}")
    elif a.cmd == "skeleton":
        print(skeleton(a.pack_id, a.role, a.force))
    elif a.cmd == "reextract":
        print(reextract(a.pack_id))
    elif a.cmd == "draft":
        def _prog(i, n, name):
            print(f"  [{i}/{n}] {name}", flush=True)
        print(draft(a.pack_id, a.force, llm=a.llm, model=a.model,
                    progress=(_prog if a.llm else None), resume=a.resume,
                    only=(a.only.split(",") if a.only else None),
                    max_cv_chars=a.max_chars))
    elif a.cmd == "build":
        rep = build(a.pack_id, a.dest, a.no_deploy, a.skip_verify, a.skip_weak)
        print("deployed:", len(rep["deployed"]), "files")
        print("docs copied:", rep["docs_copied"])
        if rep["missing_refs"]:
            print("missing refs:")
            for n, f in rep["missing_refs"]:
                print("  ", n, "->", f)
        if rep["problems"]:
            print("warnings:")
            for p_ in rep["problems"]:
                print("  ", p_)
        if "verify" in rep:
            for e in rep["verify"]["candidates"]:
                print(("OK  " if e["ok"] else "FAIL") + "  " + e["candidate"],
                      "|", "; ".join(e["checks"]))
    elif a.cmd == "verify":
        mf = load_manifest(a.pack_id)
        dest = a.dest or os.path.join(DEST_ROOT, mf["pack_name"])
        v = verify_pack(dest, mf["pack_name"])
        for e in v["candidates"]:
            print(("OK  " if e["ok"] else "FAIL") + "  " + e["candidate"],
                  "|", "; ".join(e["checks"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())